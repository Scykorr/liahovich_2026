from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from charset_normalizer import from_bytes
from django.conf import settings
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

ALLOWED_EXTENSIONS = {".csv", ".xlsx"}
CSV_SIGNATURE_PRINTABLE = True
XLSX_MAGIC = b"PK"
CANONICAL_FIELDS = [
    "territory_code",
    "precinct_code",
    "registered_voters",
    "ballots_in_boxes",
    "outside_ballots",
    "valid_ballots",
    "invalid_ballots",
]
CANONICAL_LABELS = {
    "territory_code": "Код территории (ТИК)",
    "precinct_code": "Код УИК",
    "registered_voters": "Избиратели",
    "ballots_in_boxes": "Бюллетени в ящиках",
    "outside_ballots": "Вне участка",
    "valid_ballots": "Действительные",
    "invalid_ballots": "Недействительные",
}


def sanitize_filename(name: str) -> str:
    base = Path(name).name
    base = base.replace("\x00", "")
    if base in {"", ".", ".."} or "/" in base or "\\" in base:
        raise ValidationError("Некорректное имя файла")
    return base


def validate_upload(file_obj, filename: str) -> None:
    filename = sanitize_filename(filename)
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValidationError("Допустимы только файлы .csv и .xlsx")
    size = getattr(file_obj, "size", None)
    if size is not None and size > settings.MAX_UPLOAD_SIZE_BYTES:
        raise ValidationError(
            f"Файл больше допустимого размера ({settings.MAX_UPLOAD_SIZE_BYTES // (1024 * 1024)} МБ)"
        )
    position = file_obj.tell()
    header = file_obj.read(8)
    file_obj.seek(position)
    if suffix == ".xlsx" and not header.startswith(XLSX_MAGIC):
        raise ValidationError("Содержимое файла не соответствует XLSX")
    content_type = getattr(file_obj, "content_type", "") or ""
    allowed_types = {
        "text/csv",
        "application/vnd.ms-excel",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
        "text/plain",
        "",
    }
    if content_type and content_type.split(";")[0].strip() not in allowed_types:
        raise ValidationError("Недопустимый MIME-тип файла")


def detect_csv_meta(raw: bytes) -> tuple[str, str]:
    charset = from_bytes(raw).best()
    encoding = charset.encoding if charset else "utf-8"
    if encoding.lower() in {"utf-8", "utf_8"} and raw.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    text = raw.decode(encoding, errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t,")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    return encoding, delimiter


def read_tabular(path: Path, encoding: str | None = None, delimiter: str | None = None) -> tuple[list[str], list[dict[str, str]], str, str]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValidationError("Пустая таблица")
        columns = [str(cell).strip() if cell is not None else f"col_{i}" for i, cell in enumerate(header_row)]
        rows: list[dict[str, str]] = []
        for raw in rows_iter:
            if all(cell is None or str(cell).strip() == "" for cell in raw):
                continue
            rows.append(
                {
                    columns[i]: "" if i >= len(raw) or raw[i] is None else str(raw[i]).strip()
                    for i in range(len(columns))
                }
            )
        workbook.close()
        return columns, rows, "xlsx", ""
    raw = path.read_bytes()
    detected_encoding, detected_delimiter = detect_csv_meta(raw)
    encoding = encoding or detected_encoding
    delimiter = delimiter or detected_delimiter
    text = raw.decode(encoding, errors="replace")
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    columns = [name or f"col_{i}" for i, name in enumerate(reader.fieldnames or [])]
    rows = []
    for item in reader:
        rows.append({key: ("" if value is None else str(value).strip()) for key, value in item.items()})
    return columns, rows, encoding, delimiter


def guess_mapping(columns: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    normalized = {col: _norm(col) for col in columns}
    aliases = {
        "territory_code": ("territory", "tik", "тик", "территория", "код_территории"),
        "precinct_code": ("precinct", "uik", "уик", "участок", "код_уик"),
        "registered_voters": ("registered", "voters", "избирател", "список"),
        "ballots_in_boxes": ("ballots_in_boxes", "ящик", "выдано", "бюллетен"),
        "outside_ballots": ("outside", "вне", "надом"),
        "valid_ballots": ("valid", "действительн"),
        "invalid_ballots": ("invalid", "недействительн"),
    }
    used: set[str] = set()
    for field, keys in aliases.items():
        for column, norm in normalized.items():
            if column in used:
                continue
            if any(key in norm for key in keys):
                mapping[field] = column
                used.add(column)
                break
    return mapping


def guess_candidates(columns: list[str], mapping: dict[str, str]) -> list[str]:
    used = set(mapping.values())
    result = []
    for column in columns:
        if column in used:
            continue
        norm = _norm(column)
        if norm.startswith("candidate") or "кандидат" in norm:
            result.append(column)
    if len(result) < 2:
        result = [column for column in columns if column not in used]
    return result


def _norm(value: str) -> str:
    return re.sub(r"[^0-9a-zа-яё]+", "_", value.lower()).strip("_")
